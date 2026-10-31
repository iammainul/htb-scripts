#!/usr/bin/env python3
"""
extract_fields.py - pull chosen fields out of one or more files and save them as CSV or TXT.

Supported inputs (picked by file extension):
  .xlsx / .xlsm             -> first non-empty row is the header; fields match header names
  .json / .jsonl / .ndjson  -> fields can be nested paths, e.g. user.name or items.0.id
  .csv / .tsv               -> fields are column headers
  .xml                      -> fields are tag names or XPath-ish paths (e.g. ./host/ip)
  anything else (.txt, .log, .conf ...) -> lines like "key: value" or "key = value"

Usage examples:
  python extract_fields.py -f a.json b.json -k name email address.city -o out.csv
  python extract_fields.py -f data.xlsx -k name email -o out.csv
  python extract_fields.py -f *.log -k username ip -o out.txt
  python extract_fields.py -f data.csv -k id status --format txt
  python extract_fields.py            # no args -> interactive prompts
"""
import argparse, csv, datetime, glob, json, os, re, sys
import xml.etree.ElementTree as ET


def get_path(obj, path):
    """Walk a dotted path through dicts/lists. Returns None if missing."""
    cur = obj
    for part in path.split("."):
        if isinstance(cur, dict):
            cur = cur.get(part)
        elif isinstance(cur, list) and part.isdigit() and int(part) < len(cur):
            cur = cur[int(part)]
        else:
            return None
        if cur is None:
            return None
    if isinstance(cur, (dict, list)):
        return json.dumps(cur, ensure_ascii=False)
    return cur


def records_json(path):
    with open(path, encoding="utf-8") as fh:
        if path.lower().endswith((".jsonl", ".ndjson")):
            return [json.loads(l) for l in fh if l.strip()]
        data = json.load(fh)
    return data if isinstance(data, list) else [data]


def records_csv(path):
    delim = "\t" if path.lower().endswith(".tsv") else ","
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh, delimiter=delim))


def extract_xml(path, fields):
    root = ET.parse(path).getroot()
    row = {}
    for f in fields:
        q = f if f.startswith(".") else ".//" + f
        hits = root.findall(q)
        row[f] = "; ".join((h.text or "").strip() for h in hits) if hits else None
    return [row]


KV = re.compile(r"^\s*([^:=#]+?)\s*[:=]\s*(.*?)\s*$")

def extract_text(path, fields):
    wanted = {f.lower(): f for f in fields}
    row = {f: None for f in fields}
    with open(path, encoding="utf-8", errors="replace") as fh:
        for line in fh:
            m = KV.match(line)
            if m and m.group(1).lower() in wanted:
                key = wanted[m.group(1).lower()]
                row[key] = m.group(2) if row[key] is None else f"{row[key]}; {m.group(2)}"
    return [row]


def _cell_value(v):
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return v.isoformat()
    return v


def _row_empty(values):
    for v in values:
        if v is None:
            continue
        if isinstance(v, str) and not v.strip():
            continue
        return False
    return True


def extract_xlsx(path, fields, sheets=None):
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("Reading .xlsx/.xlsm requires openpyxl. Install it with: pip install openpyxl")

    wanted = {f.strip().lower(): f for f in fields}
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        available = {n.lower(): n for n in wb.sheetnames}
        if sheets:
            selected = []
            for name in sheets:
                if name in wb.sheetnames:
                    selected.append(name)
                elif name.lower() in available:
                    selected.append(available[name.lower()])
                else:
                    print(f"[skip] {path}: sheet {name!r} not found", file=sys.stderr)
        else:
            selected = list(wb.sheetnames)

        out = []
        for name in selected:
            header = None
            for raw in wb[name].iter_rows(values_only=True):
                if _row_empty(raw):
                    continue
                if header is None:
                    header = {}
                    seen = set()
                    for i, cell in enumerate(raw):
                        if cell is None:
                            continue
                        key = str(cell).strip().lower()
                        if key in wanted and wanted[key] not in seen:
                            header[i] = wanted[key]
                            seen.add(wanted[key])
                    continue
                row = {f: None for f in fields}
                for i, cell in enumerate(raw):
                    if i in header:
                        row[header[i]] = _cell_value(cell)
                row["sheet"] = name
                out.append(row)
        return out
    finally:
        wb.close()


def extract(path, fields, sheets=None):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return extract_xlsx(path, fields, sheets)
    if ext in (".json", ".jsonl", ".ndjson"):
        return [{f: get_path(r, f) for f in fields} for r in records_json(path)]
    if ext in (".csv", ".tsv"):
        return [{f: r.get(f) for f in fields} for r in records_csv(path)]
    if ext == ".xml":
        return extract_xml(path, fields)
    return extract_text(path, fields)


def expand(patterns):
    out = []
    for p in patterns:
        matches = glob.glob(p)
        out.extend(matches if matches else [p])
    return out


def main():
    ap = argparse.ArgumentParser(description="Extract fields from files into CSV/TXT.")
    ap.add_argument("-f", "--files", nargs="+", help="input files (wildcards ok)")
    ap.add_argument("-k", "--fields", nargs="+", help="fields to extract")
    ap.add_argument("-o", "--output", default="extracted.csv", help="output file (.csv or .txt)")
    ap.add_argument("--format", choices=["csv", "txt"], help="force output format")
    ap.add_argument("--sheet", nargs="+", help="xlsx/xlsm sheet names (default: all sheets)")
    ap.add_argument("--no-source", action="store_true", help="omit the source_file column")
    a = ap.parse_args()

    if not a.files:
        a.files = input("File names (space or comma separated): ").replace(",", " ").split()
    if not a.fields:
        a.fields = input("Fields to extract (space or comma separated): ").replace(",", " ").split()
    if not a.files or not a.fields:
        sys.exit("Need at least one file and one field.")

    paths = expand(a.files)
    fmt = a.format or ("txt" if a.output.lower().endswith(".txt") else "csv")
    xlsx = any(os.path.splitext(p)[1].lower() in (".xlsx", ".xlsm") for p in paths)
    cols = ([] if a.no_source else ["source_file"]) + (["sheet"] if xlsx else []) + a.fields
    rows = []
    for path in paths:
        try:
            for r in extract(path, a.fields, a.sheet):
                if not a.no_source:
                    r = {"source_file": os.path.basename(path), **r}
                rows.append(r)
        except Exception as e:
            print(f"[skip] {path}: {e}", file=sys.stderr)

    with open(a.output, "w", encoding="utf-8", newline="") as fh:
        if fmt == "csv":
            w = csv.DictWriter(fh, fieldnames=cols)
            w.writeheader()
            for r in rows:
                w.writerow({c: "" if r.get(c) is None else r.get(c) for c in cols})
        else:
            fh.write("\t".join(cols) + "\n")
            for r in rows:
                fh.write("\t".join("" if r.get(c) is None else str(r.get(c)) for c in cols) + "\n")

    print(f"Wrote {len(rows)} row(s) to {a.output}")


if __name__ == "__main__":
    main()
